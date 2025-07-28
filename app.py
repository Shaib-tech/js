from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, flash
import pandas as pd
import os
from datetime import datetime, timedelta, time
import json
from io import BytesIO
import xlsxwriter
from openpyxl import load_workbook
from werkzeug.utils import secure_filename
import calendar

app = Flask(__name__)
app.secret_key = 'your-secret-key-here'

# Configuration
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'xlsx', 'xls'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Ensure upload directory exists
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs('data', exist_ok=True)

class AttendanceManager:
    def __init__(self):
        self.employees_file = 'data/employees.json'
        self.attendance_file = 'data/attendance.json'
        self.salary_file = 'data/salary.json'
        self.load_data()
    
    def load_data(self):
        """Load data from JSON files"""
        try:
            with open(self.employees_file, 'r') as f:
                self.employees = json.load(f)
        except FileNotFoundError:
            self.employees = {}
        
        try:
            with open(self.attendance_file, 'r') as f:
                self.attendance = json.load(f)
        except FileNotFoundError:
            self.attendance = {}
        
        try:
            with open(self.salary_file, 'r') as f:
                self.salary_data = json.load(f)
        except FileNotFoundError:
            self.salary_data = {}
    
    def save_data(self):
        """Save data to JSON files"""
        with open(self.employees_file, 'w') as f:
            json.dump(self.employees, f, indent=2)
        
        with open(self.attendance_file, 'w') as f:
            json.dump(self.attendance, f, indent=2)
        
        with open(self.salary_file, 'w') as f:
            json.dump(self.salary_data, f, indent=2)
    
    def add_employee(self, emp_id, name, shift_type, monthly_salary):
        """Add a new employee"""
        self.employees[emp_id] = {
            'name': name,
            'shift_type': shift_type,
            'monthly_salary': float(monthly_salary)
        }
        self.save_data()
    
    def mark_attendance(self, emp_id, date, arrival_time):
        """Mark attendance for an employee"""
        if emp_id not in self.employees:
            return False, "Employee not found"
        
        employee = self.employees[emp_id]
        shift_type = employee['shift_type']
        
        # Parse arrival time
        try:
            arrival = datetime.strptime(arrival_time, '%H:%M').time()
        except ValueError:
            return False, "Invalid time format"
        
        # Determine attendance status based on shift
        status = self.calculate_attendance_status(shift_type, arrival)
        
        # Store attendance
        if date not in self.attendance:
            self.attendance[date] = {}
        
        self.attendance[date][emp_id] = {
            'arrival_time': arrival_time,
            'status': status
        }
        
        self.save_data()
        return True, f"Attendance marked as {status}"
    
    def calculate_attendance_status(self, shift_type, arrival_time):
        """Calculate attendance status based on shift and arrival time"""
        if shift_type == 'First':
            punctual_start = time(10, 0)
            punctual_end = time(10, 15)
            grace_end = time(10, 30)
            half_day_start = time(11, 0)
        else:  # Second shift
            punctual_start = time(14, 0)
            punctual_end = time(14, 15)
            grace_end = time(14, 30)
            half_day_start = time(15, 0)
        
        if punctual_start <= arrival_time <= punctual_end:
            return 'Punctual'
        elif arrival_time <= grace_end:
            return 'Grace'
        elif arrival_time >= half_day_start:
            return 'Half-day'
        else:
            return 'Late'
    
    def get_monthly_summary(self, emp_id, year, month):
        """Get monthly attendance summary for an employee"""
        summary = {
            'punctual': 0,
            'grace': 0,
            'late': 0,
            'half_day': 0,
            'absent': 0,
            'total_working_days': 0
        }
        
        # Get all days in the month
        _, last_day = calendar.monthrange(year, month)
        working_days = 0
        
        for day in range(1, last_day + 1):
            date = f"{year}-{month:02d}-{day:02d}"
            weekday = datetime(year, month, day).weekday()
            
            # Assuming Monday-Friday are working days (0-4)
            if weekday < 5:
                working_days += 1
                if date in self.attendance and emp_id in self.attendance[date]:
                    status = self.attendance[date][emp_id]['status']
                    summary[status.lower().replace('-', '_')] += 1
                else:
                    summary['absent'] += 1
        
        summary['total_working_days'] = working_days
        return summary
    
    def calculate_salary(self, emp_id, year, month, bonus=0):
        """Calculate salary for an employee for a specific month"""
        if emp_id not in self.employees:
            return None
        
        employee = self.employees[emp_id]
        monthly_salary = employee['monthly_salary']
        summary = self.get_monthly_summary(emp_id, year, month)
        
        # Calculate per day salary
        per_day_salary = monthly_salary / summary['total_working_days']
        
        # Calculate deductions
        absent_deduction = summary['absent'] * per_day_salary
        half_day_deduction = summary['half_day'] * (per_day_salary * 0.5)
        
        total_deductions = absent_deduction + half_day_deduction
        net_salary = monthly_salary - total_deductions + bonus
        
        salary_details = {
            'employee_id': emp_id,
            'employee_name': employee['name'],
            'shift_type': employee['shift_type'],
            'basic_salary': monthly_salary,
            'per_day_salary': per_day_salary,
            'attendance_summary': summary,
            'absent_deduction': absent_deduction,
            'half_day_deduction': half_day_deduction,
            'total_deductions': total_deductions,
            'bonus': bonus,
            'net_salary': net_salary,
            'year': year,
            'month': month
        }
        
        # Store salary data
        month_key = f"{year}-{month:02d}"
        if month_key not in self.salary_data:
            self.salary_data[month_key] = {}
        
        self.salary_data[month_key][emp_id] = salary_details
        self.save_data()
        
        return salary_details

# Initialize the attendance manager
manager = AttendanceManager()

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/')
def dashboard():
    return render_template('dashboard.html', employees=manager.employees)

@app.route('/employees')
def employees():
    return render_template('employees.html', employees=manager.employees)

@app.route('/add_employee', methods=['POST'])
def add_employee():
    data = request.get_json()
    emp_id = data['emp_id']
    name = data['name']
    shift_type = data['shift_type']
    monthly_salary = data['monthly_salary']
    
    if emp_id in manager.employees:
        return jsonify({'success': False, 'message': 'Employee ID already exists'})
    
    manager.add_employee(emp_id, name, shift_type, monthly_salary)
    return jsonify({'success': True, 'message': 'Employee added successfully'})

@app.route('/attendance')
def attendance():
    return render_template('attendance.html', employees=manager.employees)

@app.route('/mark_attendance', methods=['POST'])
def mark_attendance():
    data = request.get_json()
    emp_id = data['emp_id']
    date = data['date']
    arrival_time = data['arrival_time']
    
    success, message = manager.mark_attendance(emp_id, date, arrival_time)
    return jsonify({'success': success, 'message': message})

@app.route('/salary')
def salary():
    return render_template('salary.html', employees=manager.employees)

@app.route('/calculate_salary', methods=['POST'])
def calculate_salary():
    data = request.get_json()
    emp_id = data['emp_id']
    year = int(data['year'])
    month = int(data['month'])
    bonus = float(data.get('bonus', 0))
    
    salary_details = manager.calculate_salary(emp_id, year, month, bonus)
    if salary_details:
        return jsonify({'success': True, 'salary_details': salary_details})
    else:
        return jsonify({'success': False, 'message': 'Employee not found'})

@app.route('/get_attendance_summary/<emp_id>/<int:year>/<int:month>')
def get_attendance_summary(emp_id, year, month):
    summary = manager.get_monthly_summary(emp_id, year, month)
    return jsonify(summary)

@app.route('/import_excel', methods=['POST'])
def import_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file selected'})
    
    file = request.files['file']
    import_type = request.form['import_type']
    
    if file.filename == '':
        return jsonify({'success': False, 'message': 'No file selected'})
    
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        try:
            if import_type == 'employees':
                df = pd.read_excel(filepath)
                required_columns = ['Employee ID', 'Employee Name', 'Shift Type', 'Monthly Salary']
                
                if not all(col in df.columns for col in required_columns):
                    return jsonify({'success': False, 'message': 'Missing required columns'})
                
                for _, row in df.iterrows():
                    manager.add_employee(
                        str(row['Employee ID']),
                        row['Employee Name'],
                        row['Shift Type'],
                        row['Monthly Salary']
                    )
                
                return jsonify({'success': True, 'message': f'Imported {len(df)} employees'})
            
            elif import_type == 'attendance':
                df = pd.read_excel(filepath)
                required_columns = ['Employee ID', 'Date', 'Arrival Time']
                
                if not all(col in df.columns for col in required_columns):
                    return jsonify({'success': False, 'message': 'Missing required columns'})
                
                imported_count = 0
                for _, row in df.iterrows():
                    date = pd.to_datetime(row['Date']).strftime('%Y-%m-%d')
                    arrival_time = str(row['Arrival Time'])
                    if ':' not in arrival_time:
                        # Handle time as decimal (e.g., 10.5 = 10:30)
                        hours = int(float(arrival_time))
                        minutes = int((float(arrival_time) - hours) * 60)
                        arrival_time = f"{hours:02d}:{minutes:02d}"
                    
                    success, _ = manager.mark_attendance(str(row['Employee ID']), date, arrival_time)
                    if success:
                        imported_count += 1
                
                return jsonify({'success': True, 'message': f'Imported {imported_count} attendance records'})
        
        except Exception as e:
            return jsonify({'success': False, 'message': f'Error processing file: {str(e)}'})
        finally:
            if os.path.exists(filepath):
                os.remove(filepath)
    
    return jsonify({'success': False, 'message': 'Invalid file format'})

@app.route('/export_excel/<export_type>')
def export_excel(export_type):
    output = BytesIO()
    
    if export_type == 'employees':
        with xlsxwriter.Workbook(output) as workbook:
            worksheet = workbook.add_worksheet('Employees')
            
            # Headers
            headers = ['Employee ID', 'Employee Name', 'Shift Type', 'Monthly Salary']
            for col, header in enumerate(headers):
                worksheet.write(0, col, header)
            
            # Data
            row = 1
            for emp_id, emp_data in manager.employees.items():
                worksheet.write(row, 0, emp_id)
                worksheet.write(row, 1, emp_data['name'])
                worksheet.write(row, 2, emp_data['shift_type'])
                worksheet.write(row, 3, emp_data['monthly_salary'])
                row += 1
        
        filename = 'employees_export.xlsx'
    
    elif export_type == 'attendance':
        with xlsxwriter.Workbook(output) as workbook:
            worksheet = workbook.add_worksheet('Attendance')
            
            # Headers
            headers = ['Date', 'Employee ID', 'Employee Name', 'Shift Type', 'Arrival Time', 'Status']
            for col, header in enumerate(headers):
                worksheet.write(0, col, header)
            
            # Data
            row = 1
            for date, daily_attendance in manager.attendance.items():
                for emp_id, att_data in daily_attendance.items():
                    if emp_id in manager.employees:
                        emp_data = manager.employees[emp_id]
                        worksheet.write(row, 0, date)
                        worksheet.write(row, 1, emp_id)
                        worksheet.write(row, 2, emp_data['name'])
                        worksheet.write(row, 3, emp_data['shift_type'])
                        worksheet.write(row, 4, att_data['arrival_time'])
                        worksheet.write(row, 5, att_data['status'])
                        row += 1
        
        filename = 'attendance_export.xlsx'
    
    elif export_type == 'salary':
        with xlsxwriter.Workbook(output) as workbook:
            worksheet = workbook.add_worksheet('Salary Summary')
            
            # Headers
            headers = ['Month', 'Employee ID', 'Employee Name', 'Shift Type', 'Basic Salary', 
                      'Present Days', 'Half Days', 'Absent Days', 'Total Deductions', 'Bonus', 'Net Salary']
            for col, header in enumerate(headers):
                worksheet.write(0, col, header)
            
            # Data
            row = 1
            for month_key, month_data in manager.salary_data.items():
                for emp_id, salary_data in month_data.items():
                    worksheet.write(row, 0, month_key)
                    worksheet.write(row, 1, emp_id)
                    worksheet.write(row, 2, salary_data['employee_name'])
                    worksheet.write(row, 3, salary_data['shift_type'])
                    worksheet.write(row, 4, salary_data['basic_salary'])
                    
                    summary = salary_data['attendance_summary']
                    present_days = summary['punctual'] + summary['grace'] + summary['late']
                    worksheet.write(row, 5, present_days)
                    worksheet.write(row, 6, summary['half_day'])
                    worksheet.write(row, 7, summary['absent'])
                    worksheet.write(row, 8, salary_data['total_deductions'])
                    worksheet.write(row, 9, salary_data['bonus'])
                    worksheet.write(row, 10, salary_data['net_salary'])
                    row += 1
        
        filename = 'salary_export.xlsx'
    
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/generate_salary_slip/<emp_id>/<int:year>/<int:month>')
def generate_salary_slip(emp_id, year, month):
    month_key = f"{year}-{month:02d}"
    
    if month_key not in manager.salary_data or emp_id not in manager.salary_data[month_key]:
        return jsonify({'success': False, 'message': 'Salary data not found'})
    
    salary_data = manager.salary_data[month_key][emp_id]
    return render_template('salary_slip.html', salary_data=salary_data)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)